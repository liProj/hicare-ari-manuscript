"""Population denominators for the ARI rates of Medicina 2026, 62, 1752.

Source: INEC, "Estimaciones y Proyecciones de la Poblacion de Ecuador, Revision 2024" (based on
the 2022 census), tabulados Nacional.zip and Provincial.zip from
https://www.ecuadorencifras.gob.ec/proyecciones-poblacionales/ .  The files on the site are the
third edition (December 2025) of that revision; values are mid-year (30 June) populations.
WHO world standard population (2000-2025): https://seer.cancer.gov/stdpopulations/world.who.html

Writes, into data/pop/:
  pop_national_age.csv   year, sex, age_group, population      (21 five-year groups, 0-4 .. 100+)
  pop_province.csv       year, prov_code, province, population, pop_0_19, pop_20plus
  pop_province_age.csv   year, prov_code, province, sex, age_group, population
  who_standard.csv       age_group, weight_pct, seer_per_million

Every sheet is checked: the age groups must add up to the printed Total row.
"""
import html
import os
import re

import openpyxl
import pandas as pd

JD = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
POP = f"{JD}/data/pop"
RAW = f"{POP}/raw"
YEARS = range(2015, 2024)

AGES = [f"{a}-{a + 4}" for a in range(0, 100, 5)] + ["100+"]
# sheet order of the provincial workbook == INEC province code order
PROV = [("01", "Azuay", "azuay"), ("02", "Bolívar", "bolivar"), ("03", "Cañar", "cañar"),
        ("04", "Carchi", "carchi"), ("05", "Cotopaxi", "cotopaxi"),
        ("06", "Chimborazo", "chimborazo"), ("07", "El Oro", "el_oro"),
        ("08", "Esmeraldas", "esmeraldas"), ("09", "Guayas", "guayas"),
        ("10", "Imbabura", "imbabura"), ("11", "Loja", "loja"), ("12", "Los Ríos", "los_rios"),
        ("13", "Manabí", "manabi"), ("14", "Morona Santiago", "morona"), ("15", "Napo", "napo"),
        ("16", "Pastaza", "pastaza"), ("17", "Pichincha", "pichincha"),
        ("18", "Tungurahua", "tungurahua"), ("19", "Zamora Chinchipe", "zamora"),
        ("20", "Galápagos", "galapagos"), ("21", "Sucumbíos", "sucumbios"),
        ("22", "Orellana", "orellana"), ("23", "Santo Domingo de los Tsáchilas", "santo_domingo"),
        ("24", "Santa Elena", "santa_elena")]
SEX = {"both": ("población_ambos_sexos", "n"), "male": ("población_hombres", "h"),
       "female": ("población_mujeres", "m")}


def norm_age(label):
    s = re.sub(r"\s+", "", str(label))
    if s.lower().startswith("100"):
        return "100+"
    return s if s in AGES else None


def parse_sheet(ws):
    """-> DataFrame(year, age_group, population) for YEARS, checked against the Total row."""
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(r for r in rows if sum(isinstance(v, int) and 1900 < v < 2100 for v in r) > 20)
    col = {v: i for i, v in enumerate(hdr) if isinstance(v, int) and v in YEARS}
    assert len(col) == len(YEARS), col
    rec, total = [], None
    for r in rows:
        lab = r[1]
        if lab is None:
            continue
        if str(lab).strip() == "Total":
            total = {y: r[i] for y, i in col.items()}
            continue
        a = norm_age(lab)
        if a:
            rec += [dict(year=y, age_group=a, population=int(round(r[i]))) for y, i in col.items()]
    df = pd.DataFrame(rec)
    assert df.age_group.nunique() == len(AGES) and len(df) == len(AGES) * len(YEARS)
    s = df.groupby("year").population.sum()
    for y in YEARS:      # rounding of the published integers only
        assert abs(s[y] - total[y]) <= len(AGES), (ws.title, y, s[y], total[y])
    return df


def national():
    wb = openpyxl.load_workbook(f"{RAW}/Nacional/tabul_nac_edad_quin_1950-2050.xlsx",
                                read_only=True, data_only=True)
    parts = [parse_sheet(wb[sheet]).assign(sex=sex) for sex, (sheet, _) in SEX.items()]
    return pd.concat(parts)[["year", "sex", "age_group", "population"]]


def provincial():
    wb = openpyxl.load_workbook(
        f"{RAW}/Provincial/Tabulado_provincial_edad_quinquenal_1990-2035.xlsx",
        read_only=True, data_only=True)
    parts = []
    for code, name, key in PROV:
        for sex, (_, suf) in SEX.items():
            parts.append(parse_sheet(wb[f"{key}_{suf}"]).assign(
                prov_code=code, province=name, sex=sex))
    return pd.concat(parts)[["year", "prov_code", "province", "sex", "age_group", "population"]]


def who():
    h = open(f"{RAW}/who/world.who.html", errors="ignore").read()
    rec = []
    for tr in re.findall(r"<tr.*?</tr>", h, re.S):
        c = [html.unescape(re.sub(r"<[^>]+>", "", x)).strip()
             for x in re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", tr, re.S)]
        if len(c) == 5 and c[0] in AGES:
            rec.append(dict(age_group=c[0], weight_pct=float(c[1]),
                            seer_per_million=int(c[4].replace(",", "").rstrip("*"))))
    df = pd.DataFrame(rec)
    assert list(df.age_group) == AGES and df.seer_per_million.sum() == 1_000_000
    return df


def main():
    nat = national()
    pa = provincial()
    w = who()
    young = set(AGES[:4])
    b = pa[pa.sex == "both"]
    prov = (b.groupby(["year", "prov_code", "province"]).population.sum().reset_index()
            .merge(b[b.age_group.isin(young)].groupby(["year", "prov_code"]).population.sum()
                   .rename("pop_0_19").reset_index(), on=["year", "prov_code"]))
    prov["pop_20plus"] = prov.population - prov.pop_0_19
    nat.to_csv(f"{POP}/pop_national_age.csv", index=False)
    prov.to_csv(f"{POP}/pop_province.csv", index=False)
    pa.to_csv(f"{POP}/pop_province_age.csv", index=False)
    w.to_csv(f"{POP}/who_standard.csv", index=False)

    # ---- sanity checks against the denominators implied by the article -------------------
    tot = nat[nat.sex == "both"].groupby("year").population.sum()
    mf = nat[nat.sex != "both"].groupby("year").population.sum()
    psum = prov.groupby("year").population.sum()
    print("year  national   male+female-both  sum(provinces)-national")
    for y in YEARS:
        print(y, f"{tot[y]:,}", mf[y] - tot[y], psum[y] - tot[y])
    for y, n, rate in ((2015, 48025, 295.19), (2023, 66270, 371.58)):
        implied = n / rate * 1e5
        print(f"{y}: article implies {implied:,.0f}; INEC rev. 2024 gives {tot[y]:,} "
              f"({(tot[y] / implied - 1) * 100:+.4f}%); rate from our denominator "
              f"{n / tot[y] * 1e5:.2f} vs printed {rate}")
    ari = f"{JD}/data/cohort/ari.parquet"
    if os.path.exists(ari):
        a = pd.read_parquet(ari, columns=["anio_ingr", "prov_ubi", "prov_res", "causa3"])
        a = a[(a.anio_ingr == 2015)]
        az = int(prov[(prov.year == 2015) & (prov.prov_code == "01")].population.iloc[0])
        for c in ("prov_ubi", "prov_res"):
            n = int((a[c] == "01").sum())
            print(f"Azuay 2015 by {c}: {n} ARI / {az:,} = {n / az * 1e5:.2f} (printed 454.79)")


if __name__ == "__main__":
    main()
